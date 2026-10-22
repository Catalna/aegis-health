# -*- coding: utf-8 -*-
"""Build deck proposal J-ARI — Healthkathon 2026 (18 slide, 16:9).

Semua angka hasil evaluasi diambil dari backend/models_trained/*.json.
Angka publik diberi sumber di slide masing-masing.
"""
from deck_kit import (Deck, M, W, NAVY, TEAL, INK, GRAY, LGRAY, HAIR, LINE,
                      TINT, TINT_WARM, WHITE, TITLE_FONT, BODY_FONT)
from pptx.enum.text import PP_ALIGN
from pptx.util import Pt, Inches
from pptx.dml.color import RGBColor

D = Deck()
TOTAL = 18

# ================================================================
# 01 — COVER
# ================================================================
D.new_slide()
D.text(M, 0.40, 9.5, 0.26, "HEALTHKATHON 2026 · EFISIENSI RISIKO FASILITAS KESEHATAN",
       10.5, TEAL, bold=True, font=TITLE_FONT, spacing=160)
D.text(9.2, 0.40, W - 9.2 - M, 0.26, "PROPOSAL SOLUSI", 10.5, LGRAY, bold=True,
       font=TITLE_FONT, spacing=160, align=PP_ALIGN.RIGHT)
D.hairline(M, 0.80, W - 2 * M)

D.text(M, 1.70, 8.0, 1.15, "J-ARI", 76, NAVY, bold=True, font=TITLE_FONT)
D.text(M, 2.95, 10.5, 0.45, "JKN Agentic Risk Intelligence", 22, TEAL, bold=True,
       font=TITLE_FONT)
D.text(M, 3.55, 11.6, 0.35, "Detect the risk. Verify with evidence. Protect JKN.",
       14, INK, font=BODY_FONT, italic=True)
D.text(M, 4.15, 11.4, 0.75,
       "Lapisan screening & risk intelligence yang membantu verifier BPJS Kesehatan "
       "memprioritaskan klaim JKN yang paling perlu diperiksa lebih dahulu — "
       "disertai bukti yang dapat ditelusuri dan skor yang sepenuhnya deterministik.",
       12, GRAY, line_spacing=1.25)

fy = 5.55
D.chip(M, fy, 3.30, 0.52, "Fokus: Upcoding & Unbundling", accent=True, size=10.5)
D.chip(M + 3.50, fy, 4.55, 0.52,
       "Sinyal pendukung: cloning · repeat · phantom · inflated", size=9.5)
D.chip(M + 8.25, fy, 3.83, 0.52, "Data sintetis ber-label · Human-in-the-loop",
       size=9.5)
D.text(M, 6.45, 11.5, 0.30, "Kategori: Efisiensi Risiko Fasilitas Kesehatan — "
       "Kemenkes RI × BPJS Kesehatan", 9.5, LGRAY, font=BODY_FONT)
D.footer(1, TOTAL)
D.save_clip = None

# ================================================================
# 02 — MASALAH & URGENSI (skala)
# ================================================================
D.new_slide("MASALAH & URGENSI",
            "Klaim tumbuh jauh lebih cepat daripada kapasitas verifikasi",
            "Program JKN terbesar di dunia dikelola dengan pemeriksaan klaim yang "
            "masih menyentuh setiap berkas secara manual.", 2, TOTAL)

cw = (W - 2 * M - 3 * 0.28) / 4
sy = 1.90
D.stat_card(M, sy, cw, 1.72, "282,7 jt",
            "peserta JKN aktif — hampir seluruh penduduk Indonesia",
            "Sumber: BPJS Kesehatan, 2025")
D.stat_card(M + (cw + 0.28), sy, cw, 1.72, "Rp 191,3 T",
            "pembayaran klaim layanan kesehatan sepanjang 2025; rasio klaim 108,27% — pengeluaran melebihi iuran",
            "Sumber: BPJS Kesehatan, 2026")
D.stat_card(M + 2 * (cw + 0.28), sy, cw, 1.72, "±Rp 6 T",
            "estimasi kerugian fraud bidang kesehatan per tahun",
            "Dirut BPJS Kesehatan, Juli 2026")
D.stat_card(M + 3 * (cw + 0.28), sy, cw, 1.72, "Rp 34 M",
            "dugaan fraud klaim di 6 RS (Sumut & Jateng) yang ditelusuri KPK",
            "KPK, 2024", accent=True)

by = 4.05
D.text(M, by, 12.0, 0.3, "Yang terjadi di lapangan", 13, NAVY, bold=True,
       font=TITLE_FONT)
lines = [
    "Verifier harus menelaah jutaan klaim per tahun dengan kedalaman yang sama — "
    "klaim berisiko tinggi bisa terlewat di antara tumpukan klaim yang sah.",
    "Pola manipulasi hanya terlihat ketika klaim dibandingkan lintas pasien, faskes, "
    "dokter, dan waktu — bukan saat berkas dibaca satu per satu.",
    "Setiap bulan deteksi berarti dana jaminan sudah keluar; setiap tuduhan yang keliru "
    "justru merugikan faskes yang legitimate. Prioritisasi harus akurat dan berbasis bukti.",
]
yy = by + 0.42
for i, ln in enumerate(lines, 1):
    D.text(M, yy, 0.35, 0.3, f"{i:02d}", 12, TEAL, bold=True, font=TITLE_FONT)
    D.text(M + 0.45, yy + 0.02, W - 2 * M - 0.5, 0.62, ln, 10.5, GRAY,
           line_spacing=1.15)
    yy += 0.78
D.text(M, 6.55, 12.0, 0.28,
       "Masalah dibingkai sebagai efisiensi & integritas sistem — bukan tuduhan "
       "terhadap faskes.", 9.5, LGRAY, italic=True)

# ================================================================
# 03 — AKAR MASALAH
# ================================================================
D.new_slide("MASALAH & URGENSI",
            "Pola risiko tidak terlihat dari satu klaim",
            "INA-CBG membayar per kelompok diagnosis — input grouper menentukan "
            "tarif, dan justru di situlah risiko manipulasi berada.", 3, TOTAL)

lx, ly = M, 1.95
D.para(lx, ly, 5.3, 3.4, [
    [("Input grouper = target manipulasi.  ", {"bold": True, "color": NAVY, "font": TITLE_FONT, "size": 10.5}),
     ("Severity level, diagnosis sekunder, dan prosedur menentukan kelompok INA-CBG dan tarifnya. Menaikkan input tanpa bukti klinis = klaim naik.", {})],
    [("Kesalahan ≠ kecurangan.  ", {"bold": True, "color": NAVY, "font": TITLE_FONT, "size": 10.5}),
     ("Banyak anomali adalah error coding; konteks klinis tidak selalu tampak dari data klaim. Sistem harus memberi prioritas, bukan vonis.", {})],
    [("Verifikasi seragam tidak skalabel.  ", {"bold": True, "color": NAVY, "font": TITLE_FONT, "size": 10.5}),
     ("Memeriksa semua klaim sama rata berarti tidak ada yang diperiksa mendalam. Sumber daya terbatas harus diarahkan ke risiko tertinggi.", {})],
], 10.5, GRAY, space_after=10, line_spacing=1.2)

rx = M + 5.75
rw = W - rx - M
D.text(rx, ly, rw, 0.3, "Tujuh pola risiko yang dicari", 12, NAVY, bold=True,
       font=TITLE_FONT)
patterns = [
    ("Upcoding / severity inflation", "severity 2→3/4 tanpa bukti; diagnosis sekunder \u201ckomplikasi\u201d fiktif", True),
    ("Unbundling", "satu paket INA-CBG dipecah menjadi beberapa klaim terpisah", True),
    ("Cloning", "klaim disalin antar pasien di faskes yang sama", False),
    ("Repeat billing", "episode yang sama ditagih ulang dalam <45 hari", False),
    ("Phantom billing", "menagih tindakan yang tidak pernah dilakukan", False),
    ("Inflated bill", "biaya jauh di atas pembanding sejenis", False),
    ("Prolonged LOS", "hari rawat > 2× batas atas rata-rata klinis", False),
]
py = ly + 0.42
for name, desc, focus in patterns:
    D.card(rx, py, rw, 0.50, accent=focus)
    if focus:
        D.text(rx + 0.14, py + 0.05, 0.85, 0.2, "FOKUS", 7.5, TEAL, bold=True,
               font=TITLE_FONT, spacing=120)
        nx = rx + 1.02
    else:
        nx = rx + 0.14
    D.text(nx, py + 0.14 if not focus else py + 0.05, 2.55, 0.4, name, 10, NAVY,
           bold=True, font=TITLE_FONT, line_spacing=0.95)
    desc_w = rw - (nx - rx) - 2.62 - 0.14
    D.text(nx + 2.62, py + 0.05 if focus else py + 0.14, desc_w, 0.42, desc,
           8.5, GRAY, line_spacing=1.0)
    py += 0.585

# ================================================================
# 04 — IDENTITAS & POSITIONING
# ================================================================
D.new_slide("IDENTITAS & POSITIONING",
            "J-ARI — lapisan intelijen di atas data klaim, bukan sistem baru",
            "J-ARI tidak menyentuh sistem pembayaran; ia mengurutkan klaim yang "
            "sudah masuk dan menjelaskan alasannya.", 4, TOTAL)

fy = 2.0
boxes = [
    ("Faskes / E-Klaim", "klaim masuk dari upstream", False),
    ("Sistem klaim BPJS", "Vclaim · SEP · e-klaim", False),
    ("J-ARI screening", "validasi · 3 sinyal · skor", True),
    ("Priority queue + evidence", "High / Medium / Low + alasan", False),
    ("Verifier memutuskan", "keputusan & audit trail", False),
]
bw, gap = 2.24, 0.28
bx = M
for i, (t, s, acc) in enumerate(boxes):
    D.card(bx, fy, bw, 0.92, accent=acc)
    D.text(bx + 0.12, fy + 0.16 if not acc else fy + 0.10, bw - 0.24, 0.4, t, 10.5,
           NAVY, bold=True, font=TITLE_FONT, line_spacing=0.95,
           align=PP_ALIGN.CENTER)
    D.text(bx + 0.10, fy + (0.52 if acc else 0.50), bw - 0.2, 0.36, s, 8, GRAY,
           align=PP_ALIGN.CENTER, line_spacing=1.0)
    if i < 4:
        D.arrow_right(bx + bw + 0.03, fy + 0.46, bx + bw + gap - 0.03)
    bx += bw + gap

D.card(M, 3.35, W - 2 * M, 1.05, accent=True)
D.text(M + 0.2, 3.50, 0.9, 0.3, "POSISI", 8.5, TEAL, bold=True, font=TITLE_FONT,
       spacing=140)
D.text(M + 0.2, 3.76, W - 2 * M - 0.4, 0.6,
       "\u201cJ-ARI membantu verifier BPJS memprioritaskan klaim JKN yang paling perlu "
       "diperiksa dengan menggabungkan sinyal klinis, reimbursement, dan pola historis "
       "ke dalam skor risiko yang transparan serta bukti yang dapat ditelusuri.\u201d",
       11.5, INK, font=BODY_FONT, line_spacing=1.2)

D.card(M, 4.62, W - 2 * M, 0.72)
D.text(M + 0.2, 4.74, 1.3, 0.3, "BATASAN", 8.5, LGRAY, bold=True, font=TITLE_FONT,
       spacing=140)
D.text(M + 0.2, 4.98, W - 2 * M - 0.4, 0.3,
       "Tidak menggantikan sistem klaim BPJS  ·  tidak menggantikan verifier  ·  "
       "tidak memvonis fraud — output tertinggi: \u201chigh risk, perlu verifikasi\u201d.",
       10, GRAY)

val = [
    ("Bagi verifier", "fokus waktu pada klaim paling berisiko, dengan alasan siap kutip"),
    ("Bagi tim investigasi", "pola antar-klaim, preseden serupa, dan jejak audit lengkap"),
    ("Bagi faskes & peserta", "klarifikasi lebih tepat sasaran; dana jaminan terlindungi"),
]
vy = 5.62
for t, s in val:
    D.text(M, vy, 2.4, 0.3, t, 10, NAVY, bold=True, font=TITLE_FONT)
    D.text(M + 2.5, vy + 0.01, 9.5, 0.3, s, 10, GRAY)
    vy += 0.40

# ================================================================
# 05 — SOLUSI: APA YANG DITERIMA VERIFIER
# ================================================================
D.new_slide("SOLUSI & KEUNGGULAN",
            "Verifier menerima antrian yang sudah diurutkan — lengkap alasannya",
            "Dalam 30 detik membaca satu halaman, verifier tahu klaim mana yang "
            "dibuka lebih dulu dan mengapa.", 5, TOTAL)

rows = [
    ("01", "Priority queue", "Klaim diurutkan High / Medium / Low memakai skor risiko "
     "deterministik — urutan tidak pernah diubah oleh LLM."),
    ("02", "Evidence terstruktur", "Setiap skor disertai bukti per sinyal (klinis, "
     "benchmark, pola) yang dapat dikutip saat klarifikasi dengan faskes."),
    ("03", "Klaim serupa", "Preseden klaim mirip (kode, biaya, faskes, waktu) sebagai "
     "alat bantu menduga cloning & repeat billing."),
    ("04", "Narasi Bahasa Indonesia", "Ringkasan investigasi + panel tanya-jawab "
     "terpandu dari LLM lokal — membaca bukti, bukan mengarang angka."),
    ("05", "Keputusan & audit trail", "Verified Valid · Klarifikasi · Perlu Review · "
     "Escalate — setiap keputusan tercatat dengan reviewer dan waktunya."),
]
yy = 1.95
for num, t, b in rows:
    D.text(M, yy + 0.03, 0.55, 0.3, num, 13, TEAL, bold=True, font=TITLE_FONT)
    D.text(M + 0.55, yy, 3.0, 0.3, t, 12, NAVY, bold=True, font=TITLE_FONT)
    D.text(M + 3.7, yy + 0.02, W - M - 3.7 - M, 0.6, b, 10, GRAY, line_spacing=1.12)
    yy += 0.72
    if num != "05":
        D.hairline(M, yy - 0.10, W - 2 * M)

D.card(M, yy + 0.02, W - 2 * M, 0.62, accent=True)
D.text(M + 0.2, yy + 0.10, W - 2 * M - 0.4, 0.45,
       [("Prinsip antarmuka:  ", {"bold": True, "color": TEAL, "font": TITLE_FONT, "size": 10.5}),
        ("skor dan alasan deterministik tampil lebih dulu; narasi AI menyusul sebagai penjelas — tidak pernah sebaliknya.", {"color": INK})],
       10.5, INK, font=BODY_FONT, anchor=None)

# ================================================================
# 06 — CARA KERJA (8 langkah)
# ================================================================
D.new_slide("SOLUSI & KEUNGGULAN",
            "Delapan langkah dari input klaim sampai keputusan verifier",
            "Validasi deterministik di depan, tiga agen paralel di tengah, "
            "skor non-LLM dan keputusan manusia di akhir.", 6, TOTAL)

CARD_W = (W - 2 * M - 3 * 0.30) / 4
ROW_Y = [1.92, 4.18]
CARD_H = 1.98
steps = [
    ("01", "Input klaim di E-Klaim", "Diagnosis, tindakan, obat, lama rawat, biaya, dan dokumen pendukung."),
    ("02", "Klaim terkirim ke BPJS", "Klaim memiliki ID sumber dan waktu penerimaan."),
    ("03", "Klaim masuk ke J-ARI", "Dikonversi ke format klaim standar; provenance dicatat."),
    ("04", "Validasi otomatis", "Data quality, validasi tanggal & field, cek duplikasi."),
    ("05", "Tiga agen menganalisis paralel", "Clinical — diagnosis, severity, LOS.  Financial — benchmark, deviasi.  Pattern — similarity, unbundling."),
    ("06", "Hasil agen jadi evidence", "Coverage, sinyal yang saling mendukung, dan uncertainty."),
    ("07", "Risk Engine menghitung prioritas", "Skor risiko deterministik, kategori High / Medium / Low, rekomendasi aksi."),
    ("08", "Verifier memutuskan", "Valid · Klarifikasi · Eskalasi — tiap keputusan masuk audit trail."),
]
for idx, (num, title, body) in enumerate(steps):
    row, col = divmod(idx, 4)
    x = M + col * (CARD_W + 0.30)
    y = ROW_Y[row]
    accent = (num == "07")
    D.card(x, y, CARD_W, CARD_H, accent=accent)
    pad = 0.16
    D.text(x + pad, y + 0.13, CARD_W - 2 * pad, 0.26, num, 13, TEAL, bold=True,
           font=TITLE_FONT, spacing=100)
    D.text(x + pad, y + 0.42, CARD_W - 2 * pad, 0.62, title, 12.5, NAVY, bold=True,
           font=TITLE_FONT, line_spacing=1.0)
    D.text(x + pad, y + 1.02, CARD_W - 2 * pad, 0.80, body, 9, GRAY,
           line_spacing=1.12)
    if accent:
        D.text(x + pad, y + CARD_H - 0.34, CARD_W - 2 * pad, 0.22,
               "DETERMINISTIK · NON-LLM", 7.5, TEAL, bold=True, font=TITLE_FONT,
               spacing=120)
for row in range(2):
    ay = ROW_Y[row] + CARD_H / 2
    for col in range(3):
        x1 = M + col * (CARD_W + 0.30) + CARD_W + 0.05
        D.arrow_right(x1, ay, x1 + 0.20)
x_end = M + 3 * (CARD_W + 0.30) + CARD_W / 2
x_start = M + CARD_W / 2
mid_y = (ROW_Y[0] + CARD_H + ROW_Y[1]) / 2
D.polyline_arrow([(x_end, ROW_Y[0] + CARD_H + 0.02), (x_end, mid_y),
                  (x_start, mid_y), (x_start, ROW_Y[1] - 0.04)])
fb_y = ROW_Y[1] + CARD_H + 0.30
D.polyline_arrow([(x_end, ROW_Y[1] + CARD_H + 0.02), (x_end, fb_y),
                  (x_start, fb_y), (x_start, ROW_Y[1] + CARD_H + 0.06)], dashed=True)
lbl = D.slide.shapes.add_textbox(Inches(4.47), Inches(fb_y - 0.13), Inches(4.4),
                                 Inches(0.26))
lbl.fill.solid(); lbl.fill.fore_color.rgb = WHITE; lbl.line.fill.background()
ltf = lbl.text_frame; ltf.word_wrap = False
ltf.margin_left = ltf.margin_right = Inches(0.08)
ltf.margin_top = ltf.margin_bottom = 0
p = ltf.paragraphs[0]; p.alignment = 1
r = p.add_run(); r.text = "Feedback keputusan verifier → evaluasi model"
r.font.name = BODY_FONT; r.font.size = Pt(9); r.font.italic = True
r.font.color.rgb = LGRAY

# ================================================================
# 07 — TIGA SINYAL
# ================================================================
D.new_slide("SOLUSI & KEUNGGULAN",
            "Tiga sinyal saling menguatkan — skor tetap murni deterministik",
            "Klaim dengan satu sinyal lemah berbeda kualitas dari klaim dengan tiga "
            "sinyal kuat sekaligus; risk engine menimbang semuanya.", 7, TOTAL)

sig = [
    ("CLINICAL", "Heuristik konsistensi INA-CBG",
     ["Severity ≥3 tanpa prosedur pendukung", "Diagnosis sekunder risiko upcoding (A41.9, Z87.39, E46)",
      "Obat & lama rawat tidak konsisten dengan severity"]),
    ("FINANCIAL", "Benchmark reimbursement peer",
     ["XGBoost memprediksi benchmark yang diharapkan", "SHAP menjelaskan fitur paling berpengaruh",
      "Deviasi >50% / >100% terhadap pembanding ditandai"]),
    ("PATTERN", "Similarity & pola relasional",
     ["LanceDB 384-d — kemiripan kode & biaya", "Skor komposit: kode 40% · biaya 30% · faskes 15% · waktu 10% · pasien 5%",
      "Window 45 hari: cloning, repeat billing, unbundling"]),
]
cw3 = (W - 2 * M - 2 * 0.30) / 3
cy = 1.92
for i, (tag, t, items) in enumerate(sig):
    x = M + i * (cw3 + 0.30)
    D.card(x, cy, cw3, 2.55)
    D.text(x + 0.16, cy + 0.14, cw3 - 0.32, 0.24, tag, 8.5, TEAL, bold=True,
           font=TITLE_FONT, spacing=140)
    D.text(x + 0.16, cy + 0.40, cw3 - 0.32, 0.5, t, 12.5, NAVY, bold=True,
           font=TITLE_FONT, line_spacing=1.0)
    D.para(x + 0.16, cy + 0.95, cw3 - 0.32, 1.5,
           ["·  " + it for it in items], 9, GRAY, line_spacing=1.12, space_after=5)

ry = 4.72
D.card(M, ry, W - 2 * M, 1.62, accent=True)
D.text(M + 0.2, ry + 0.12, 6.0, 0.26, "RISK ENGINE — DETERMINISTIK, NON-LLM", 9,
       TEAL, bold=True, font=TITLE_FONT, spacing=120)
parts = [("Clinical", "maks 20"), ("Financial", "maks 30"), ("Similarity", "maks 20"),
         ("Evidence support", "maks 15")]
px = M + 0.2
for name, mx in parts:
    D.text(px, ry + 0.48, 2.1, 0.28, name, 10.5, NAVY, bold=True, font=TITLE_FONT)
    D.text(px, ry + 0.76, 2.1, 0.28, mx + " poin", 9.5, GRAY)
    px += 2.30
D.hairline(M + 0.2, ry + 1.14, W - 2 * M - 0.4)
D.text(M + 0.2, ry + 1.24, W - 2 * M - 0.4, 0.3,
       "Isolation Forest & graph centrality dicatat sebagai konteks perilaku berbobot 0 — "
       "baru diaktifkan sebagai poin bila ablation membuktikan nilai tambahnya.",
       9, GRAY, italic=True)

# ================================================================
# 08 — PEMBEDA
# ================================================================
D.new_slide("SOLUSI & KEUNGGULAN", "Bukan fraud detector generik",
            "Keunggulan J-ARI ada pada domain knowledge, transparansi, dan tata kelola "
            "LLM — bukan pada jumlah modelnya.", 8, TOTAL)

rows = [
    ["Dimensi", "Fraud detector generik", "J-ARI"],
    ["Pengetahuan domain", "Anomali statistik generik",
     "Heuristik konsistensi INA-CBG: diagnosis + severity + prosedur → tarif"],
    ["Sinyal", "1–2 sinyal (biasanya finansial)",
     "Tiga sinyal terintegrasi: klinis, benchmark, pola relasional"],
    ["Explainability", "Skor tunggal tanpa konteks",
     "SHAP + evidence per agen + narasi Bahasa Indonesia, traceable per evidence_id"],
    ["Output", "\u201cFraud / bukan fraud\u201d",
     ("Prioritas verifikasi + bukti untuk investigasi manusia", NAVY, True)],
    ["Tata kelola LLM", "LLM sering menentukan keputusan",
     "LLM hanya menulis narasi; skor & tier selalu deterministik"],
]
# header band dulu, baru baris tabel — agar baris pertama tidak tertutup
D.card(M, 1.95, W - 2 * M, 0.34, fill=TINT_WARM, border=HAIR)
for j, htxt in enumerate(rows[0]):
    D.text(M + [0.12, 2.62, 6.82][j], 2.0, [2.3, 4.0, 5.2][j], 0.26, htxt, 9,
           NAVY, bold=True, font=TITLE_FONT)
D.table(M, 2.29, W - 2 * M, [2.5, 4.2, 5.39], None, rows[1:], row_h=0.70,
        size=10)

D.card(M, 5.90, W - 2 * M, 0.72, accent=True)
D.text(M + 0.2, 6.02, W - 2 * M - 0.4, 0.5,
       [("Yang tidak kami klaim:  ", {"bold": True, "color": TEAL, "font": TITLE_FONT, "size": 10.5}),
        ("J-ARI tidak membuktikan fraud, tidak menolak klaim otomatis, dan tidak menggantikan verifier — ia memastikan waktu verifier habis untuk klaim yang paling layak diperiksa.", {"color": INK})],
       10.5, INK)

# ================================================================
# 09 — ARSITEKTUR
# ================================================================
D.new_slide("PENDEKATAN TEKNIS & DATA", "Arsitektur ringkas — berjalan offline penuh",
            "Semua komponen open-source; LLM 7B berjalan lokal via Ollama tanpa "
            "kirim data keluar.", 9, TOTAL)

# baris 1: UI + API
b1 = [("UI Verifier", "React + Vite — 7 halaman", False),
      ("FastAPI", "REST + SSE real-time", False),
      ("Orchestrator 5 fase", "gate deterministik → 3 agen paralel → sintesis → skor → narasi", True)]
bw2, gp2 = 3.83, 0.30
bx = M
for i, (t, s, acc) in enumerate(b1):
    D.card(bx, 1.92, bw2, 0.95, accent=acc)
    D.text(bx + 0.16, 1.92 + (0.14 if acc else 0.18), bw2 - 0.32, 0.3, t, 11.5, NAVY,
           bold=True, font=TITLE_FONT, align=PP_ALIGN.CENTER)
    D.text(bx + 0.16, 2.36, bw2 - 0.32, 0.45, s, 8.5, GRAY, align=PP_ALIGN.CENTER,
           line_spacing=1.0)
    if i < 2:
        D.arrow_right(bx + bw2 + 0.03, 2.40, bx + bw2 + 0.27)
    bx += bw2 + gp2

# panah turun
D.polyline_arrow([(W / 2, 2.87), (W / 2, 3.22)])

# baris 2: storage & LLM
b2 = [("PostgreSQL", "klaim, analisis, keputusan manusia, audit log", False),
      ("LanceDB", "vektor 384-d untuk similarity search", False),
      ("Ollama — LLM lokal", "qwen2.5 7B · deepseek-r1 7B · mistral 7B (Q4_K_M)", False)]
bx = M
for t, s, acc in b2:
    D.card(bx, 3.22, bw2, 0.95)
    D.text(bx + 0.16, 3.36, bw2 - 0.32, 0.3, t, 11.5, NAVY, bold=True,
           font=TITLE_FONT, align=PP_ALIGN.CENTER)
    D.text(bx + 0.16, 3.68, bw2 - 0.32, 0.45, s, 8.5, GRAY, align=PP_ALIGN.CENTER,
           line_spacing=1.0)
    bx += bw2 + gp2

# NFR chips
nfr = ["< 3,5 dtk per klaim (termasuk LLM lokal)", "Fallback non-blocking per agen — pipeline tidak crash",
       "On-premise / offline penuh", "Cukup 1 GPU 8 GB (laptop kelas menengah)"]
cy2 = 4.62
cwid = (W - 2 * M - 3 * 0.28) / 4
for i, t in enumerate(nfr):
    D.chip(M + i * (cwid + 0.28), cy2, cwid, 0.72, t, accent=(i == 0), size=9)

D.text(M, 5.70, 12.0, 0.3, "Keputusan teknis kunci", 12, NAVY, bold=True,
       font=TITLE_FONT)
tech = [
    ("LLMClient protocol", "Ollama lokal / Gemini API dapat ditukar tanpa mengubah agen — hindari vendor lock-in."),
    ("Idempotensi & race guard", "analysis_generation UUID + trigger immutable pada received_at & queue_sequence."),
    ("Evaluasi sebagai bagian produk", "endpoint /api/ml/eval menampilkan metrik nyata — bukan angka simulasi."),
]
ty = 6.05
for t, s in tech:
    D.text(M, ty, 2.6, 0.28, t, 9.5, NAVY, bold=True, font=TITLE_FONT)
    D.text(M + 2.75, ty + 0.01, 9.3, 0.28, s, 9.5, GRAY)
    ty += 0.32

# ================================================================
# 10 — DATA & PRIVASI
# ================================================================
D.new_slide("PENDEKATAN TEKNIS & DATA",
            "Kebutuhan data: skema klaim standar — tanpa identitas pribadi",
            "MVP berjalan dengan data sintetis ber-label; produksi memakai klaim "
            "upstream melalui adapter terstandar.", 10, TOTAL)

D.text(M, 1.95, 5.6, 0.3, "Field yang dibutuhkan (sudah distandarkan)", 12, NAVY,
       bold=True, font=TITLE_FONT)
fields = [
    ("Identitas terpseudonim", "patient_id PAT-XXXXX · provider FAC-X99 · doctor DR-X99-XXX"),
    ("Diagnosis & prosedur", "ICD-10 primer + sekunder · ICD-9-CM"),
    ("Karakter episode", "severity 1–4 · lama rawat · tanggal masuk/pulang"),
    ("Finansial & farmasi", "biaya klaim · daftar obat + kuantitas"),
    ("Provenance", "ID sumber, waktu terima, nomor antrean — immutable"),
]
fy2 = 2.35
for t, s in fields:
    D.text(M, fy2, 2.4, 0.3, t, 10, NAVY, bold=True, font=TITLE_FONT,
           line_spacing=1.0)
    D.text(M + 2.55, fy2 + 0.01, 3.3, 0.5, s, 9, GRAY, line_spacing=1.05)
    fy2 += 0.62
D.text(M, fy2 + 0.05, 5.8, 0.55,
       "Sumber produksi: aliran klaim E-Klaim/Vclaim faskes — J-ARI menerima salinan "
       "read-only, tidak pernah mengubah sistem sumber.", 9.5, LGRAY, italic=True,
       line_spacing=1.15)

rx2 = M + 6.1
rw2 = W - rx2 - M
D.card(rx2, 1.95, rw2, 2.30, accent=True)
D.text(rx2 + 0.2, 2.08, rw2 - 0.4, 0.28, "MVP — DATA SINTETIS BER-LABEL", 9, TEAL,
       bold=True, font=TITLE_FONT, spacing=120)
D.para(rx2 + 0.2, 2.42, rw2 - 0.4, 1.7, [
    "·  Generator data dipisah dari pembuat detector (anti-tautologi)",
    "·  Test set held-out dengan seed berbeda — tidak dilihat saat tuning",
    "·  Ground truth disimpan terpisah dari input model",
    "·  Hasil evaluasi ditampilkan apa adanya, termasuk kegagalannya",
], 9.5, GRAY, space_after=5, line_spacing=1.1)

D.card(rx2, 4.40, rw2, 1.95)
D.text(rx2 + 0.2, 4.53, rw2 - 0.4, 0.28, "KALIBRASI & AKSES DATA RESMI BPJS", 9, NAVY,
       bold=True, font=TITLE_FONT, spacing=120)
D.para(rx2 + 0.2, 4.87, rw2 - 0.4, 1.35, [
    "·  Agregat: 100 penyakit terbanyak & kepadatan faskes per provinsi — menyelaraskan generator sintetis dengan profil JKN nyata",
    "·  Data Sampel BPJS: cohort longitudinal 2015–2024 (±126 juta baris, terde-identifikasi) via permintaan resmi",
    "·  Jalur sah ke data tingkat klaim untuk validasi — tanpa menyalahi aturan data peserta",
], 9.5, GRAY, space_after=5, line_spacing=1.1)

# ================================================================
# 11 — KEMATANGAN
# ================================================================
D.new_slide("TINGKAT KEMATANGAN", "Prototype end-to-end yang sudah berjalan",
            "Bukan konsep di atas kertas — seluruh alur dari intake sampai keputusan "
            "manusia dapat didemonstrasikan offline.", 11, TOTAL)

D.chip(M, 1.92, 1.6, 0.4, "PROTOTYPE", accent=True, size=10)
D.text(M + 1.8, 1.97, 10.5, 0.3,
       "Level 3 dari 5 — fitur utama berfungsi penuh dalam lingkungan terbatas.",
       10, GRAY)

checks = [
    ("Pipeline 5 fase lengkap", "gate deterministik → 3 agen paralel → sintesis → risk engine → narasi, dengan fallback non-blocking per agen."),
    ("UI verifier 7 halaman", "queue berprioritas, investigasi real-time (SSE), assessment SHAP + QA panel, monitoring, evaluasi ML, intake."),
    ("Keputusan & audit trail", "setiap eksekusi agen, keputusan manusia, dan reviewer tercatat (agent_runs, human_decisions, audit_logs)."),
    ("Evaluasi otomatis", "CLI & endpoint metrik: Precision@K, Recall@K, PR-AUC, latensi p50/p95 — dijalankan ulang setiap perubahan."),
    ("Offline-first", "LLM 7B lokal via Ollama; demo end-to-end tanpa koneksi eksternal; cukup 1 GPU 8 GB."),
    ("Data sintetis disiplin", "generator terpisah dari detector; test set held-out seed berbeda; ground truth terpisah."),
]
cy3 = 2.55
for i, (t, s) in enumerate(checks):
    col = i % 2
    row = i // 2
    x = M + col * (6.15)
    y = cy3 + row * 1.14
    D.text(x, y, 0.3, 0.3, "\u2713", 12, TEAL, bold=True, font=TITLE_FONT)
    D.text(x + 0.35, y, 5.4, 0.28, t, 11, NAVY, bold=True, font=TITLE_FONT)
    D.text(x + 0.35, y + 0.30, 5.5, 0.68, s, 9, GRAY, line_spacing=1.12)

D.card(M, 6.05, W - 2 * M, 0.62)
D.text(M + 0.2, 6.17, W - 2 * M - 0.4, 0.42,
       [("Batasan yang jujur:  ", {"bold": True, "color": NAVY, "font": TITLE_FONT, "size": 10}),
        ("belum diuji dengan data klaim riil — validasi domain expert dan pilot dengan data berizin masuk Fase 4.", {"color": GRAY})],
       10, GRAY)

# ================================================================
# 12 — HASIL EVALUASI
# ================================================================
D.new_slide("TINGKAT KEMATANGAN", "Hasil evaluasi nyata — diukur pada test set tersendiri",
            "Angka dari pipeline evaluasi otomatis pada 78 klaim uji profil "
            "\u201cchallenge\u201d dengan 25 anomali tanam — bukan angka simulasi.", 12, TOTAL)

cw4 = (W - 2 * M - 3 * 0.28) / 4
sy4 = 1.95
D.stat_card(M, sy4, cw4, 1.80, "100%",
            "Recall@50 — 25 dari 25 anomali berada di 50 klaim teratas",
            "Held-out challenge set, seed 2026", accent=True)
D.stat_card(M + (cw4 + 0.28), sy4, cw4, 1.80, "0,74 → 0,79",
            "Average Precision naik +5,1 pp setelah konteks perilaku ditambahkan",
            "Ablation p0-core vs +behavioral")
D.stat_card(M + 2 * (cw4 + 0.28), sy4, cw4, 1.80, "118 / 184 ms",
            "Latensi p50 / p95 per klaim pada risk engine (tanpa LLM)",
            "Pipeline evaluation runner")
D.stat_card(M + 3 * (cw4 + 0.28), sy4, cw4, 1.80, "MAPE 5,2%",
            "Akurasi benchmark biaya pada klaim normal (holdout)",
            "XGBoost regressor, 3.464 sampel")

D.text(M, 4.05, 12.0, 0.3, "Model klasifikasi biner (evaluasi terpisah, 3.464 sampel)",
       12, NAVY, bold=True, font=TITLE_FONT)
D.table(M, 4.45, W - 2 * M, [2.4, 2.3, 2.3, 2.3, 2.69],
        ["Metrik", "PR-AUC", "ROC-AUC", "Presisi", "Recall"],
        [["XGBoost classifier", "0,89", "0,91", "1,00", "0,80"]],
        row_h=0.42, size=10)

D.text(M, 5.35, 12.0, 0.3, "Bacaan yang jujur atas angka ini", 12, NAVY, bold=True,
       font=TITLE_FONT)
D.para(M, 5.72, 12.1, 1.1, [
    "·  Precision@50 = 0,50 mengikuti prevalensi anomali 32% pada set uji — pada data produksi yang jauh lebih tidak seimbang, presisi top-K justru lebih bermakna.",
    "·  Label pada data sintetis bukan pengganti validasi domain — validasi eksternal adalah target Fase 4.",
    "·  Calibration score belum dihitung (skor risiko bukan probabilitas) — dicatat sebagai keterbatasan, bukan disembunyikan.",
], 9.5, GRAY, space_after=4, line_spacing=1.1)

# ================================================================
# 13 — INSIGHT EVALUASI
# ================================================================
D.new_slide("TINGKAT KEMATANGAN", "Pelajaran penting: satu model tidak cukup",
            "Klasifikasi per-klaim melewatkan pola relasional — di situlah sinyal "
            "similarity & jendela 45 hari menutup celah.", 13, TOTAL)

rows13 = [
    ["Tipe anomali", "Classifier tunggal", "Pipeline multi-sinyal", "Peringkat teratas"],
    ["Upcoding", "100%", "100%", "#6 dari 78"],
    ["Phantom billing", "100%", "100%", "#7 dari 78"],
    ["Unbundling", "100%", "100%", "#1 dari 78"],
    ["Inflated bill", "100%", "100%", "#19 dari 78"],
    ["Prolonged LOS", "100%", "100%", "#12 dari 78"],
    [("Cloning", RGBColor(0xB3, 0x26, 0x1E), True), ("0%", RGBColor(0xB3, 0x26, 0x1E), True),
     ("100%", RGBColor(0x0B, 0x6E, 0x64), True), "#33 dari 78"],
    [("Repeat billing", RGBColor(0xB3, 0x26, 0x1E), True), ("0%", RGBColor(0xB3, 0x26, 0x1E), True),
     ("100%", RGBColor(0x0B, 0x6E, 0x64), True), "#24 dari 78"],
]
# header band dulu, baru baris tabel — agar baris pertama tidak tertutup
D.card(M, 1.95, 7.6, 0.34, fill=TINT_WARM, border=HAIR)
for j, htxt in enumerate(rows13[0]):
    D.text(M + [0.12, 2.22, 4.12, 6.02][j], 2.0, [2.0, 1.8, 1.8, 1.5][j], 0.26,
           htxt, 8.5, NAVY, bold=True, font=TITLE_FONT)
D.table(M, 2.29, 7.6, [2.1, 1.9, 1.9, 1.7], None, rows13[1:], row_h=0.46, size=9.5)

rx3 = M + 8.0
rw3 = W - rx3 - M
D.para(rx3, 2.0, rw3, 3.2, [
    [("Cloning & repeat billing mustahil dilihat dari satu klaim.  ", {"bold": True, "color": NAVY, "font": TITLE_FONT, "size": 10.5}),
     ("Keduanya pola antar-klaim: pasien berbeda dengan kode & biaya nyaris identik di faskes sama, atau episode sama ditagih ulang.", {})],
    [("Sinyal relasional menutup celah.  ", {"bold": True, "color": NAVY, "font": TITLE_FONT, "size": 10.5}),
     ("Similarity berbasis vektor + aturan jendela 45 hari mengangkat 25/25 anomali ke top-50.", {})],
    [("Konteks perilaku diperlakukan jujur.  ", {"bold": True, "color": NAVY, "font": TITLE_FONT, "size": 10.5}),
     ("Isolation Forest menaikkan AP +5,1 pp, namun tetap berbobot 0 sampai ablation pada dataset lebih besar membuktikan nilainya.", {})],
], 10, GRAY, space_after=10, line_spacing=1.18)

D.card(M, 5.85, W - 2 * M, 0.72)
D.text(M + 0.2, 5.97, W - 2 * M - 0.4, 0.5,
       [("Implikasi desain:  ", {"bold": True, "color": NAVY, "font": TITLE_FONT, "size": 10.5}),
        ("arsitektur multi-sinyal bukan pemolesan — ia menjawab kegagalan spesifik yang terukur pada evaluasi kami sendiri.", {"color": INK})],
       10.5, INK)

# ================================================================
# 14 — DAMPAK & NILAI
# ================================================================
D.new_slide("DAMPAK & NILAI", "Dampak yang diusulkan — terukur, tanpa overclaim",
            "Semua perubahan dibandingkan terhadap kondisi verifikasi manual saat ini; "
            "angka final tetap berasal dari keputusan verifier.", 14, TOTAL)

ba = [
    ("Semua klaim ditinjau dengan kedalaman sama",
     "Verifier fokus ke top-K klaim paling berisiko — skor + alasan siap sebelum berkas dibuka"),
    ("Bukti tersebar di beberapa sistem & berkas",
     "Satu halaman investigasi: SHAP, klaim serupa, narasi Bahasa Indonesia, QA panel"),
    ("Kebocoran umumnya baru ketahuan saat audit",
     "Potensi exposure terhitung dini: \u03a3 maks(0, klaim − benchmark atas) untuk tier High & Medium"),
]
yy = 1.95
for i, (before, after) in enumerate(ba, 1):
    D.card(M, yy, 5.55, 0.95, fill=TINT_WARM, border=HAIR)
    D.text(M + 0.15, yy + 0.09, 5.2, 0.26, "TANPA J-ARI", 7.5, LGRAY, bold=True,
           font=TITLE_FONT, spacing=120)
    D.text(M + 0.15, yy + 0.33, 5.25, 0.55, before, 9.5, GRAY, line_spacing=1.05)
    D.arrow_right(M + 5.68, yy + 0.47, M + 6.02, color=TEAL, weight=1.25)
    D.card(M + 6.15, yy, 6.02, 0.95, accent=True)
    D.text(M + 6.30, yy + 0.09, 5.7, 0.26, "DENGAN J-ARI", 7.5, TEAL, bold=True,
           font=TITLE_FONT, spacing=120)
    D.text(M + 6.30, yy + 0.33, 5.75, 0.55, after, 9.5, INK, line_spacing=1.05)
    yy += 1.13

D.text(M, yy + 0.08, 12.0, 0.3, "Penerima manfaat", 12, NAVY, bold=True,
       font=TITLE_FONT)
bm = [
    ("Verifier BPJS", "waktu fokus & keputusan konsisten"),
    ("Tim investigasi", "pola sistemik & preseden siap pakai"),
    ("Faskes legitimate", "klarifikasi tepat sasaran, bukan tuduhan"),
    ("Peserta JKN", "dana jaminan terjaga"),
]
bw3 = (W - 2 * M - 3 * 0.28) / 4
by3 = yy + 0.44
for i, (t, s) in enumerate(bm):
    D.card(M + i * (bw3 + 0.28), by3, bw3, 0.78)
    D.text(M + i * (bw3 + 0.28) + 0.15, by3 + 0.10, bw3 - 0.3, 0.26, t, 10, NAVY,
           bold=True, font=TITLE_FONT)
    D.text(M + i * (bw3 + 0.28) + 0.15, by3 + 0.38, bw3 - 0.3, 0.34, s, 8.5, GRAY,
           line_spacing=1.0)

# ================================================================
# 15 — RISIKO, PRIVASI & ETIKA
# ================================================================
D.new_slide("RISIKO, PRIVASI & ETIKA", "Dipikirkan sejak awal, bukan dicatut belakangan",
            "Empat keputusan tata kelola yang menentukan desain sistem.", 15, TOTAL)

et = [
    ("PERLINDUNGAN DATA (UU PDP)",
     ["Tanpa nama — patient/provider/doctor berupa pseudonim",
      "Minimisasi PII: hanya field yang dibutuhkan deteksi",
      "Deploy on-premise BPJS; data tidak keluar; semua akses tercatat"]),
    ("MANUSIA DALAM KEPUTUSAN",
     ["Skor = rekomendasi prioritas, bukan vonis fraud",
      "Jawaban \u201capakah ini fraud?\u201d dibakukan di sistem",
      "Keputusan akhir selalu verifier — tercatat dengan identitasnya"]),
    ("BIAS & KETERBATASAN AI",
     ["Model belajar dari data sintetis — validasi domain ditunda ke pilot",
      "Evidence ditampilkan terbuka agar keputusan bisa diuji & dikoreksi",
      "Konteks perilaku berbobot 0 sampai ablation membuktikan nilainya"]),
    ("MITIGASI RISIKO UTAMA",
     ["False positive → jalur klarifikasi, bukan penolakan otomatis",
      "Halusinasi LLM → LLM hanya menulis narasi; tak bisa mengubah skor",
      "Drift → versioning analisis + monitoring + retraining dari keputusan verifier"]),
]
cw2x = (W - 2 * M - 0.30) / 2
for i, (tag, items) in enumerate(et):
    col, row = i % 2, i // 2
    x = M + col * (cw2x + 0.30)
    y = 1.95 + row * 2.32
    D.card(x, y, cw2x, 2.14)
    D.text(x + 0.18, y + 0.13, cw2x - 0.36, 0.26, tag, 9, TEAL, bold=True,
           font=TITLE_FONT, spacing=120)
    D.para(x + 0.18, y + 0.50, cw2x - 0.36, 1.5, ["·  " + it for it in items],
           9.5, GRAY, space_after=6, line_spacing=1.12)

D.text(M, 6.55, 12.0, 0.28,
       "Data peserta JKN riil tidak digunakan tanpa izin resmi — sesuai ketentuan "
       "Healthkathon 2026.", 9.5, LGRAY, italic=True)

# ================================================================
# 16 — RENCANA & KELAYAKAN
# ================================================================
D.new_slide("RENCANA & KELAYAKAN", "Empat fase — tiga pertama sudah dijalankan",
            "Rencana pasca-hackathon realistis: validasi domain, pilot berizin, lalu "
            "integrasi terstandar.", 16, TOTAL)

phases = [
    ("FASE 1", "Konsistensi domain & produk", "SELESAI",
     ["Fokus tajam: upcoding & unbundling", "Alur dimulai dari incoming claims",
      "Benchmark = pembanding peer, bukan \u201charga benar\u201d"]),
    ("FASE 2", "Validasi model", "SELESAI",
     ["Blind test held-out dengan seed berbeda", "Ablation: kontribusi tiap sinyal",
      "Precision@K · Recall@K · PR-AUC · latensi"]),
    ("FASE 3", "Demo & tata kelola", "SELESAI",
     ["Demo end-to-end offline satu alur", "Audit trail & disclaimer lengkap",
      "Batasan sistem didokumentasikan"]),
    ("FASE 4", "Pasca-hackathon", "RENCANA",
     ["Validasi bersama domain expert BPJS", "Pilot dengan data berizin",
      "Adapter SATUSEHAT/FHIR → feedback loop retraining"]),
]
pw = (W - 2 * M - 3 * 0.28) / 4
py2 = 1.95
for i, (ph, t, status, items) in enumerate(phases):
    x = M + i * (pw + 0.28)
    acc = status == "SELESAI"
    D.card(x, py2, pw, 3.15, accent=acc)
    D.text(x + 0.16, py2 + 0.14, pw - 0.32, 0.24, ph, 8.5,
           TEAL if acc else LGRAY, bold=True, font=TITLE_FONT, spacing=140)
    D.text(x + 0.16, py2 + 0.40, pw - 0.32, 0.6, t, 11.5, NAVY, bold=True,
           font=TITLE_FONT, line_spacing=1.0)
    D.text(x + 0.16, py2 + 0.98, pw - 0.32, 0.24, status, 8,
           TEAL if acc else LGRAY, bold=True, font=TITLE_FONT, spacing=120)
    D.para(x + 0.16, py2 + 1.32, pw - 0.32, 1.7, ["·  " + it for it in items],
           9, GRAY, space_after=5, line_spacing=1.1)
    if i < 3:
        D.arrow_right(x + pw + 0.03, py2 + 1.55, x + pw + 0.25)

D.card(M, 5.45, W - 2 * M, 1.10)
D.text(M + 0.2, 5.58, 5.0, 0.26, "SUMBER DAYA YANG DIBUTUHKAN", 8.5, NAVY,
       bold=True, font=TITLE_FONT, spacing=120)
res = [("Tim", "1–3 orang — ML, domain klaim JKN, full-stack"),
       ("Teknologi", "stack open-source, tanpa lisensi berkelanjutan"),
       ("Infrastruktur", "1 GPU 8 GB untuk pilot kecil"),
       ("Prasyarat kunci", "izin data & mitra domain BPJS")]
pos = [0.2, 3.45, 6.55, 9.30]
wid = [3.05, 2.9, 2.55, 3.0]
for (t, s), px2, wd in zip(res, pos, wid):
    D.text(M + px2, 5.90, wd, 0.26, t, 9.5, NAVY, bold=True, font=TITLE_FONT)
    D.text(M + px2, 6.14, wd, 0.4, s, 8.5, GRAY, line_spacing=1.0)

# ================================================================
# 17 — TIM
# ================================================================
D.new_slide("PROFIL TIM", "Tim yang membangun J-ARI",
            "Komposisi keahlian yang menutupi tiga hal sekaligus: ML engineering, "
            "domain klaim JKN, dan full-stack development.", 17, TOTAL)

team = [
    ("Ryu Adhi Nugroho", "Lead · ML & System Architecture",
     "Merancang pipeline multi-agent, risk engine deterministik, dan kerangka evaluasi anti-tautologi.", True),
    ("[Nama Anggota 2]", "[Peran — mis. Domain & Product]",
     "[Keahlian: pemahaman alur klaim JKN/INA-CBG, product design, kebutuhan verifier.]", False),
    ("[Nama Anggota 3]", "[Peran — mis. Full-stack & Data]",
     "[Keahlian: frontend React, backend FastAPI, pipeline data & evaluasi.]", False),
]
tw = (W - 2 * M - 2 * 0.30) / 3
for i, (name, role, desc, lead) in enumerate(team):
    x = M + i * (tw + 0.30)
    D.card(x, 1.95, tw, 2.30, accent=lead)
    D.text(x + 0.18, 2.10, tw - 0.36, 0.35, name, 13, NAVY, bold=True,
           font=TITLE_FONT)
    D.text(x + 0.18, 2.48, tw - 0.36, 0.3, role, 9.5, TEAL if lead else GRAY,
           bold=True)
    D.text(x + 0.18, 2.85, tw - 0.36, 1.25, desc, 9.5, GRAY, line_spacing=1.15)

D.text(M, 4.60, 12.0, 0.3, "Mengapa tim ini mampu mengeksekusi", 12, NAVY,
       bold=True, font=TITLE_FONT)
proofs = [
    ("Dibangun dalam satu iterasi disiplin", "PRD, arsitektur, logika bisnis, dan evaluasi terdokumentasi serta konsisten dengan kode yang berjalan."),
    ("Evaluasi bukan pelengkap", "Blind test, ablation, dan metrik dijalankan otomatis — hasil apa adanya dipublikasikan di dalam produk."),
    ("Domain dipelajari secara serius", "Logika INA-CBG, pola upcoding, dan alur verifier diterjemahkan langsung menjadi aturan & evidence sistem."),
]
py3 = 5.0
for t, s in proofs:
    D.text(M, py3, 3.4, 0.28, t, 10, NAVY, bold=True, font=TITLE_FONT)
    D.text(M + 3.6, py3 + 0.01, 8.5, 0.42, s, 9.5, GRAY, line_spacing=1.05)
    py3 += 0.52

# ================================================================
# 18 — PENUTUP
# ================================================================
D.new_slide("PENUTUP", "Deteksi yang lebih cerdas, JKN yang lebih terlindungi",
            None, 18, TOTAL)

D.text(M, 2.0, 11.8, 0.8,
       "J-ARI tidak menentukan fraud dan tidak menggantikan verifier.",
       20, NAVY, bold=True, font=TITLE_FONT, line_spacing=1.1)
D.text(M, 2.85, 11.8, 0.5,
       "J-ARI memastikan klaim yang paling perlu diperiksa diperiksa lebih dahulu — "
       "dan pemeriksa itu punya alasannya.",
       20, TEAL, bold=True, font=TITLE_FONT, line_spacing=1.1)

D.hairline(M, 4.05, W - 2 * M)
vals = [
    ("Deterministik", "Skor & tier dihitung murni matematis — LLM tidak pernah menyentuhnya."),
    ("Evidence-first", "Setiap prioritas bisa ditelusuri ke aturan atau output model."),
    ("Human-led", "AI membantu memahami; keputusan tetap milik verifier."),
    ("Data-responsible", "Data sintetis/berizin, pseudonim, on-premise, teraudit."),
]
vw = (W - 2 * M - 3 * 0.28) / 4
for i, (t, s) in enumerate(vals):
    x = M + i * (vw + 0.28)
    D.text(x, 4.30, vw, 0.3, t, 11.5, NAVY, bold=True, font=TITLE_FONT)
    D.text(x, 4.62, vw, 0.85, s, 9.5, GRAY, line_spacing=1.15)

D.card(M, 5.75, W - 2 * M, 0.85, accent=True)
D.text(M + 0.2, 5.90, 2.2, 0.3, "LANGKAH BERIKUT", 9, TEAL, bold=True,
       font=TITLE_FONT, spacing=140)
D.text(M + 0.2, 6.16, W - 2 * M - 0.4, 0.35,
       "Pilot terbuka bersama verifier BPJS Kesehatan — dengan data berizin, "
       "ukuran waktu review sebelum/sesudah, dan validasi domain expert.",
       11, INK)

D.save(r"C:\Users\Radith\E-Jak\deck_redesign\jari_proposal.pptx")
print("saved 18 slides")
