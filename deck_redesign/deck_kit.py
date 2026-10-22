# -*- coding: utf-8 -*-
"""deck_kit — design system bersama untuk deck J-ARI (Healthkathon 2026).

Gaya: editorial flat. Putih, navy #003171 (Arial bold), teal #007A70,
hairline #E4E1D8, body Inter. Tanpa gradien, tanpa ikon warna-warni.
"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.oxml.ns import qn

NAVY = RGBColor(0x00, 0x31, 0x71)
TEAL = RGBColor(0x00, 0x7A, 0x70)
INK = RGBColor(0x2B, 0x2A, 0x27)
GRAY = RGBColor(0x55, 0x52, 0x4B)
LGRAY = RGBColor(0x8A, 0x87, 0x80)
HAIR = RGBColor(0xE4, 0xE1, 0xD8)
LINE = RGBColor(0xB5, 0xB2, 0xAA)
TINT = RGBColor(0xEF, 0xF6, 0xF5)
TINT_WARM = RGBColor(0xF8, 0xF6, 0xF0)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

W, H = 13.333, 7.5
M = 0.62  # margin kiri/kanan

TITLE_FONT = "Arial"
BODY_FONT = "Inter"


class Deck:
    def __init__(self):
        self.prs = Presentation()
        self.prs.slide_width = Inches(W)
        self.prs.slide_height = Inches(H)
        self.blank = self.prs.slide_layouts[6]
        self.slide = None

    # ---------- slide lifecycle ----------
    def new_slide(self, kicker=None, title=None, subtitle=None, page=None, total=18):
        self.slide = self.prs.slides.add_slide(self.blank)
        if kicker:
            self.text(M, 0.40, 9.0, 0.26, kicker, 10.5, TEAL, bold=True,
                      font=TITLE_FONT, spacing=160)
            self.text(8.2, 0.40, W - 8.2 - M, 0.26, "J-ARI", 10.5, LGRAY, bold=True,
                      font=TITLE_FONT, spacing=160, align=PP_ALIGN.RIGHT)
        if title:
            self.text(M, 0.72, W - 2 * M, 0.55, title, 26, NAVY, bold=True,
                      font=TITLE_FONT)
        if subtitle:
            self.text(M, 1.26, W - 2 * M - 0.2, 0.32, subtitle, 11.5, GRAY,
                      font=BODY_FONT)
        if kicker:
            self.hairline(M, 1.66, W - 2 * M)
        if page:
            self.footer(page, total)
        return self.slide

    def footer(self, page, total=18):
        self.hairline(M, 7.02, W - 2 * M)
        self.text(M, 7.10, 7.0, 0.24, "J-ARI · JKN AGENTIC RISK INTELLIGENCE",
                  8.5, LGRAY, font=TITLE_FONT, spacing=140)
        self.text(7.2, 7.10, W - 7.2 - M, 0.24,
                  f"HEALTHKATHON 2026 · {page:02d} / {total}", 8.5, LGRAY,
                  font=TITLE_FONT, spacing=140, align=PP_ALIGN.RIGHT)

    def save(self, path):
        self.prs.save(path)

    # ---------- primitives ----------
    def text(self, x, y, w, h, runs, size, color, bold=False,
             align=PP_ALIGN.LEFT, font=BODY_FONT, spacing=None,
             anchor=MSO_ANCHOR.TOP, line_spacing=None, italic=False):
        tb = self.slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = anchor
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.alignment = align
        if line_spacing:
            p.line_spacing = line_spacing
        if isinstance(runs, str):
            runs = [(runs, {})]
        for t, ov in runs:
            r = p.add_run()
            r.text = t
            r.font.name = ov.get("font", font)
            r.font.size = Pt(ov.get("size", size))
            r.font.bold = ov.get("bold", bold)
            r.font.italic = ov.get("italic", italic)
            r.font.color.rgb = ov.get("color", color)
            spc = ov.get("spacing", spacing)
            if spc:
                r.font._rPr.set("spc", str(spc))
        return tb

    def para(self, x, y, w, h, lines, size, color, font=BODY_FONT,
             line_spacing=1.15, space_after=None, align=PP_ALIGN.LEFT, bold=False):
        """lines: list of str or (text, overrides) or list-of-runs."""
        tb = self.slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        for i, ln in enumerate(lines):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = align
            p.line_spacing = line_spacing
            if space_after is not None:
                p.space_after = Pt(space_after)
            runs = ln if isinstance(ln, list) else [(ln, {})]
            if isinstance(ln, tuple):
                runs = [ln]
            for t, ov in runs:
                r = p.add_run()
                r.text = t
                r.font.name = ov.get("font", font)
                r.font.size = Pt(ov.get("size", size))
                r.font.bold = ov.get("bold", bold)
                r.font.italic = ov.get("italic", False)
                r.font.color.rgb = ov.get("color", color)
                spc = ov.get("spacing")
                if spc:
                    r.font._rPr.set("spc", str(spc))
        return tb

    def hairline(self, x, y, w, color=HAIR, weight=0.75):
        ln = self.slide.shapes.add_connector(
            MSO_CONNECTOR.STRAIGHT, Inches(x), Inches(y), Inches(x + w), Inches(y))
        ln.line.color.rgb = color
        ln.line.width = Pt(weight)
        return ln

    def card(self, x, y, w, h, accent=False, fill=None, border=None):
        box = self.slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y),
                                          Inches(w), Inches(h))
        box.fill.solid()
        box.fill.fore_color.rgb = fill or (TINT if accent else WHITE)
        box.line.color.rgb = border or (TEAL if accent else HAIR)
        box.line.width = Pt(1.25 if accent else 1.0)
        box.shadow.inherit = False
        style_el = box._element.find(qn("p:style"))
        if style_el is not None:
            box._element.remove(style_el)
        return box

    def chip(self, x, y, w, h, label, accent=False, size=9.5, color=None):
        c = self.card(x, y, w, h, accent=accent)
        self.text(x, y, w, h, label, size, color or (TEAL if accent else GRAY),
                  bold=accent, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        return c

    def arrow_right(self, x1, y, x2, color=LINE, weight=1.0):
        ln = self.slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1),
                                             Inches(y), Inches(x2), Inches(y))
        ln.line.color.rgb = color
        ln.line.width = Pt(weight)
        ln_el = ln.line._get_or_add_ln()
        ln_el.append(ln_el.makeelement(qn("a:tailEnd"),
                                       {"type": "triangle", "w": "med", "len": "med"}))
        return ln

    def polyline_arrow(self, points, color=LINE, weight=1.0, dashed=False):
        for i in range(len(points) - 1):
            (x1, y1), (x2, y2) = points[i], points[i + 1]
            ln = self.slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1),
                                                 Inches(y1), Inches(x2), Inches(y2))
            ln.line.color.rgb = color
            ln.line.width = Pt(weight)
            if dashed:
                ln.line.dash_style = MSO_LINE_DASH_STYLE.DASH
            if i == len(points) - 2:
                ln_el = ln.line._get_or_add_ln()
                ln_el.append(ln_el.makeelement(
                    qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"}))

    # ---------- composites ----------
    def stat_card(self, x, y, w, h, number, label, source=None, accent=False):
        self.card(x, y, w, h, accent=accent)
        pad = 0.18
        self.text(x + pad, y + 0.16, w - 2 * pad, 0.55, number, 27,
                  TEAL if accent else NAVY, bold=True, font=TITLE_FONT)
        self.text(x + pad, y + 0.78, w - 2 * pad, 0.60, label, 9.5, GRAY,
                  line_spacing=1.1)
        if source:
            self.text(x + pad, y + h - 0.36, w - 2 * pad, 0.28, source, 7.5, LGRAY,
                      italic=True, line_spacing=1.0)

    def num_row(self, x, y, w, num, title, body, num_w=0.55):
        self.text(x, y + 0.02, num_w, 0.3, num, 13, TEAL, bold=True, font=TITLE_FONT)
        self.text(x + num_w, y, 3.1, 0.55, title, 12, NAVY, bold=True, font=TITLE_FONT,
                  line_spacing=1.0)
        self.text(x + num_w + 3.25, y + 0.02, w - num_w - 3.25, 0.6, body, 9.5, GRAY,
                  line_spacing=1.1)

    def table(self, x, y, w, col_widths, header, rows, row_h=0.42,
              header_fill=None, size=9.5, bold_first_col=True):
        """Simple flat table: hairline rows, optional header band."""
        cx = [x]
        for cw in col_widths:
            cx.append(cx[-1] + cw)
        yy = y
        if header:
            self.card(x, yy, w, 0.34, fill=header_fill or TINT_WARM, border=HAIR)
            for j, htxt in enumerate(header):
                self.text(cx[j] + 0.12, yy + 0.045, col_widths[j] - 0.2, 0.26, htxt,
                          8.5, NAVY, bold=True, font=TITLE_FONT)
            yy += 0.34
        for i, row in enumerate(rows):
            for j, cell in enumerate(row):
                bold = bold_first_col and j == 0
                color = NAVY if (bold_first_col and j == 0) else GRAY
                if isinstance(cell, tuple):  # (text, color, bold)
                    cell, color, bold = cell
                self.text(cx[j] + 0.12, yy + 0.07, col_widths[j] - 0.2, row_h - 0.1,
                          cell, size, color, bold=bold, font=BODY_FONT,
                          line_spacing=1.0)
            yy += row_h
            if i < len(rows) - 1:
                self.hairline(x, yy, w)
        return yy
